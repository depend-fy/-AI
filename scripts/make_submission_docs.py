"""生成参赛文档：队伍成员声明 + 作品说明与技术方案（DOCX）。

格式参考同赛事其他队伍的提交材料，内容为本项目真实信息。
正文用宋体、标题用黑体，并显式设置 w:eastAsia（否则中文可能不按预期字体渲染）。
"""

from __future__ import annotations

import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = pathlib.Path(__file__).resolve().parents[1] / "docs" / "submission"
OUT.mkdir(parents=True, exist_ok=True)

TEAM = "旭日伴粽同行小队"
CODE = "U071"
GROUP = "本科生组"
LEADER = "宗佳雨"
MEMBER = "郑子旭"




#: A4（21cm）减去左右各 2.5cm 页边距 = 16cm ≈ 453.5pt。
#: 所有表格列宽合计必须小于这个值，否则页面会被撑宽、内容被裁切。
TEXT_WIDTH_PT = 453.0


def setup_page(doc: Document) -> None:
    """统一页面为 A4 并设置页边距（中文公文与学术文档的通用规格）。"""
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)


def set_font(run, *, name_cn: str, size: int, bold: bool = False, name_en: str = "Times New Roman"):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name_cn)
    rfonts.set(qn("w:ascii"), name_en)
    rfonts.set(qn("w:hAnsi"), name_en)


def _add_rich(paragraph, text: str, *, name_cn: str, size: int) -> None:
    """写入文本，并把 `**加粗**` 语法转成真正的粗体 run（DOCX 不解析 Markdown）。"""
    for index, part in enumerate(text.split("**")):
        if not part:
            continue
        set_font(
            paragraph.add_run(part),
            name_cn=name_cn,
            size=size,
            bold=(index % 2 == 1),
        )


def title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), name_cn="黑体", size=18, bold=True)
    p.paragraph_format.space_after = Pt(6)


def subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), name_cn="宋体", size=10.5)
    p.paragraph_format.space_after = Pt(14)


def heading(doc: Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    set_font(p.add_run(text), name_cn="黑体", size=size, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def body(doc: Document, text: str, *, indent: bool = True, size: int = 11) -> None:
    p = doc.add_paragraph()
    _add_rich(p, text, name_cn="宋体", size=size)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)


def bullet(doc: Document, text: str, size: int = 11) -> None:
    p = doc.add_paragraph(style="List Bullet")
    _add_rich(p, text, name_cn="宋体", size=size)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(2)


def make_table(doc: Document, header: list[str], rows: list[list[str]], widths: list[float]):
    total = sum(widths)
    assert total <= TEXT_WIDTH_PT, (
        f"表格列宽合计 {total}pt 超出正文宽度 {TEXT_WIDTH_PT}pt，会被页面裁切"
    )
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(text), name_cn="黑体", size=10.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(text), name_cn="宋体", size=10)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Pt(width)
    return table


# ==========================================================================
# 文档一：队伍成员声明
# ==========================================================================
decl = Document()
setup_page(decl)
title(decl, f"{TEAM} 队伍成员声明")
subtitle(decl, f"参赛编号 {CODE}　|　{GROUP}　|　成员 2 人")

heading(decl, "一、成员名单与分工")
body(
    decl,
    "本作品由两名跨学科成员协同完成。成员分别来自金融工程与人工智能专业，"
    "具体名单及分工如下：",
)
make_table(
    decl,
    ["姓名", "学校与专业", "主要分工"],
    [
        [
            f"{MEMBER}",
            "中国矿业大学　人工智能",
            "负责前期调试：开发环境搭建、模块联调与缺陷定位；检索链路与模型接入的调试；"
            "流水线各环节的输入输出核对，确保主张拆分、检索、裁决、修正、复核五步可完整跑通。",
        ],
        [
            f"{LEADER}（队长）",
            "江苏师范大学　金融工程",
            "负责后期优化与文件撰写：核验口径与误报取舍的确定；界面与文案优化；"
            "评测体系设计（标注集、三基线对照、指标口径）；实验报告、参赛材料与演示视频文稿撰写。",
        ],
    ],
    widths=[64, 104, 280],
)

heading(decl, "二、独立完成声明")
body(
    decl,
    "本人郑重声明：所提交的项目成果（包括但不限于项目说明、技术方案、源代码、实验数据、"
    "汇报材料等）是本队伍独立完成、集成与实验验证的成果。除文中已经注明引用的内容外，"
    "本作品不包含任何其他个人或集体已经发表或撰写过的作品成果。"
    "本作品使用的第三方组件、公开数据接口与开发工具的来源及使用方式，"
    "详见随附的《第三方来源清单》。",
)
body(
    decl,
    "本队同时如实声明：开发过程中使用了 AI 编程助手（DeepSeek Harness）辅助编写代码与文档；"
    "需求确认、技术路线选择、方案取舍与验收标准由队伍成员决定，全部产出经队伍成员运行验证后才被接受。"
    "此项已在《第三方来源清单》第五节中主动披露。",
)

heading(decl, "三、关于分工协作的补充说明")
body(
    decl,
    "本项目采用“方案设计—系统实现—验证评测”的协作方式，两位成员在各自专业领域内独立承担关键任务，"
    "并通过高频沟通实现跨学科融合。结合本作品在问题价值与学科融合、任务完成度与结果正确性、"
    "模型执行闭环、技术深度与系统实现、验证严谨性与异常处理、创新性与发展价值等维度的综合要求，"
    "具体协作机制如下：",
)

heading(decl, "1. 跨学科融合的协作逻辑", size=12)
body(
    decl,
    "本作品的核心难题是：如何在“宁可漏判、不可冤枉”的前提下判断一条事实主张是否成立。"
    "这一问题的解决依赖两个专业视角的互补。"
    f"{LEADER}从金融工程视角出发，将风险度量中的“误报代价高于漏报代价”“举证责任”与“可审计”"
    "等原则引入核验口径的设计：确立“无证据不判错”的判定底线，"
    "把误报率作为最核心的诚实指标（而非只看精确率），并要求每条结论都能回溯到具体证据与规则编号，"
    "使判定结果具备可复核性。"
    f"{MEMBER}依托人工智能专业的系统实现能力，负责将方案落地为可运行系统："
    "完成开发环境搭建与模块联调，调试文献检索链路与大模型接入，"
    "核对流水线各环节的输入输出，确保主张拆分、外部检索、规则裁决、修正建议与独立复核五步能够完整跑通。"
    "两者在“证据如何才算充分”这一核心判据上形成交汇，"
    "共同构成“可追溯、可解释、可复核”的事实核验机制。",
)

heading(decl, "2. 贡献对等性说明", size=12)
body(
    decl,
    f"双方贡献对等且不可替代。{MEMBER}负责前期调试，主导了系统跑通的工程工作，"
    "在技术深度与系统实现、异常处理方面承担关键任务。"
    f"{LEADER}负责后期优化与文件撰写，主导了核验口径的确定、评测体系设计与全部文档材料的撰写，"
    "在问题价值挖掘、验证严谨性与创新性提炼方面发挥重要作用。"
    "在标注集建设、三基线对照实验与失败案例分析中，双方共同完成实验设计、结果讨论与结论确认，"
    "确保实验结论的可靠性与可复现性。",
)

heading(decl, "3. 协作机制", size=12)
body(
    decl,
    "考虑到制作周期紧凑，双方未采用固定周例会形式，而是采用“阶段对齐 + 即时沟通”的轻量化协作方式："
    "阶段对齐用于确认当日目标与技术路线，即时沟通用于解决调试异常与文档细节问题。"
    "双方共同完成了核验口径的确定、标注测试集建设、三基线对照实验与失败案例分析，"
    "并在系统实现、验证严谨性、异常处理与创新性提炼等方面协同推进。"
    "以上分工内容真实有效，双方对项目成果享有同等贡献权。",
)

decl.save(str(OUT / f"{TEAM}-队伍成员声明.docx"))
print("已生成：队伍成员声明.docx")

# ==========================================================================
# 文档二：作品说明与技术方案
# ==========================================================================
doc = Document()
setup_page(doc)
title(doc, "智鉴 AI（VeriMind AI）")
subtitle(
    doc,
    "面向科研文本的事实核验 Agent：任务规划、工具调用、结果验证与异常处理的完整闭环",
)
subtitle(doc, f"队伍名称：{TEAM}　|　编号 {CODE}　|　{GROUP}　|　宗佳雨　郑子旭")

heading(doc, "摘要")
body(
    doc,
    "AI 生成或辅助撰写的科研文本中，最危险的错误不是明显胡说，而是“看起来很像真的”："
    "引用一个并不存在的 DOI、把 2013 年发表的论文写成 2008 年、文献真实存在但没有说过那句话。"
    "这类错误会随文本一起进入论文与报告，且人工逐条核对一篇综述需要数小时至数日。"
    "本作品提出并实现了一个面向科研文本的事实核验 Agent：接收待核验文本后，"
    "自动完成主张拆分、检索策略规划、外部证据获取、可解释规则裁决、最小改动修正建议与独立复核，"
    "输出可追溯的证据链与复核结论。",
    indent=True,
)
body(
    doc,
    "本作品的核心主张是：**模型不应作为结论的来源**。大模型只产出结构化判断，"
    "判定状态由确定性规则表裁决，修正建议必须通过独立复核才允许写入正文。"
    "在 100 条人工确认的标注测试集上，本方法与“直接把文本交给同一个模型判断”相比，"
    "误报率由 44.9% 降至 2.0%，证据可追溯率由 0% 提升至 69.0%，"
    "修正成功率 79.3% 且新增错误率为 0%。",
    indent=True,
)
p = doc.add_paragraph()
set_font(
    p.add_run(
        "关键词：事实核验；证据追溯；可解释规则裁决；独立复核；科研诚信；智能体工具调用"
    ),
    name_cn="宋体",
    size=11,
    bold=True,
)

heading(doc, "一、选题依据与研究目标")
heading(doc, "（一）参考方向", size=12)
bullet(doc, "智能体与执行闭环：任务规划、工具调用、执行反馈、结果验证与异常处理的完整链路。")
bullet(doc, "科研诚信与学术写作质量保障：引用准确性、文献真实性与结论一致性的自动化检查。")
bullet(doc, "可解释裁决与独立复核：以确定性规则替代模型主观评分，并引入跨厂商模型交叉复核。")
bullet(doc, "开放学术数据接口的应用：Crossref、OpenAlex、arXiv 等公开元数据源的工具化调用。")

heading(doc, "（二）研究背景与问题来源", size=12)
body(
    doc,
    "大模型已成为科研写作的常用辅助工具，但其输出中的事实性错误具有隐蔽性。"
    "本作品把这类错误归纳为三类：一是“文献不存在”，即引用了未注册的 DOI；"
    "二是“书目信息不符”，即文献真实存在但年份、期刊等信息与注册记录不一致；"
    "三是“引用与结论不匹配”，即文献真实存在，但并未支持原文所声称的内容。",
)
body(
    doc,
    "现有做法通常是让大模型自检，但本作品的实测数据显示：直接把文本交给模型判断，"
    "误报率高达 44.9%——接近一半的正确内容会被判成有问题，且每条结论都无法追问。"
    "这说明“模型自评”既不可靠，也不可审计。因此本作品要解决的问题是："
    "如何让对文本事实性的判断变成**有证据、有规则、可复核**的过程。",
)

heading(doc, "（三）研究目标", size=12)
bullet(doc, "构建一条可执行、可验证的事实核验流水线，形成五环节完整闭环。")
bullet(doc, "确立“无证据不判错”的判定底线，区分“文献存在”“书目一致”“内容被支持”三个层面。")
bullet(doc, "以确定性规则表取代模型主观评分，使每条结论可回溯到证据与规则编号。")
bullet(doc, "引入跨厂商独立复核，确保“改了”与“改对了”是两个独立问题。")
bullet(doc, "建立标注测试集与多基线对照实验，如实公开失败案例与技术边界。")

heading(doc, "二、技术方案与实现")
heading(doc, "（一）系统架构与执行闭环", size=12)
body(
    doc,
    "系统采用“界面层—编排层—能力层—数据层”的分层结构。编排层由 11 步流水线组成，"
    "覆盖赛题要求的五个环节：",
)
make_table(
    doc,
    ["环节", "实现方式", "可验证证据"],
    [
        [
            "任务规划",
            "拆分事实主张 → 判定主张类型 → 决定检索策略（DOI 精确查询 / 书目检索 / 预印本检索 / 明确不检索并说明原因）",
            "每条主张的检索痕迹：查询词、命中数、未检索原因",
        ],
        [
            "工具调用",
            "真实调用 Crossref、OpenAlex、arXiv；带最小请求间隔、结果缓存与原始响应留档",
            "证据条目中的来源链接与原始响应留档路径；检索统计",
        ],
        [
            "执行反馈",
            "11 步流水线实时回传阶段进度；每步输入输出均为结构化数据",
            "工作台的阶段进度与阶段明细面板",
        ],
        [
            "结果验证",
            "双层验证：确定性规则表裁决（6 状态 / 8 类问题 / 严重度）＋ 跨厂商独立复核（3 项确定性检查 + 2 项语义检查）",
            "每条结论可点开证据；显示命中规则编号；复核逐项给出结论与理由",
        ],
        [
            "异常处理",
            "区分「服务端明确 404」与「网络失败」；模型调用失败降级为规则引擎并如实声明；检索失败标注「暂无法核验」；命中缓存标注数据来源",
            "报告中的适用范围与局限、降级说明、规则 N0、检索失败记录",
        ],
    ],
    widths=[58, 208, 184],
)

heading(doc, "（二）判定分层与数据契约", size=12)
body(
    doc,
    "本作品的一个关键设计是把“证据能支撑什么”变成代码级约束。证据按获取深度分级"
    "（仅元数据 / 摘要 / 全文），判定按层面分级（存在性 / 书目信息 / 内容一致性 / 文内自洽）。"
    "数据模型在构造时即强制校验：仅元数据级证据不得用于内容层面判断；"
    "判定为“存在矛盾”“部分支持”“来源冲突”时必须挂接至少一条证据；"
    "判定为“已核实”必须给出支持性证据；判定为“证据不足”必须写明还缺什么。"
    "任何违反都会导致对象无法构造，从而在系统层面杜绝“看起来有证据、其实没有”的结论。",
)
body(
    doc,
    "对于无证据判断（仅在直接问模型的对照基线中使用），数据模型要求显式标记并写明理由，"
    "报告中显著标注，不允许与有证据的结论混淆。",
)

heading(doc, "（三）修正与独立复核", size=12)
body(
    doc,
    "修正环节生成最小改动的修改方案：能改数字就不改整句，缺证据时只能弱化或标注不确定性，"
    "不得替换为新的具体数值。每条修正须通过独立复核，复核包含三项确定性检查"
    "（是否有证据支持、证据与修正后表述是否匹配、是否引入新事实错误）"
    "与两项语义检查（原始问题是否解决、是否出现新的无依据结论）。"
    "复核未通过或未完成的修正一律不写入正文。",
)
body(
    doc,
    "复核由**另一家厂商的模型**承担，以降低同源偏差；当复核模型不可用时，"
    "系统降级为规则引擎并在报告中如实声明，绝不用规则结果冒充模型判断。",
)

heading(doc, "（四）模型与调用方式", size=12)
body(doc, "本作品的模型配置如下（界面与报告均如实标注当前运行模式）：")
make_table(
    doc,
    ["角色", "模型", "职责"],
    [
        ["主模型", "deepseek-chat（DeepSeek 开放平台）", "主张拆分、事实审查、修正建议生成"],
        ["独立复核模型", "agnes-3.0-flash（Agnes）", "复核修正建议（3 项确定性 + 2 项语义检查）"],
    ],
    widths=[76, 176, 198],
)
bullet(doc, "调用协议：OpenAI 兼容的 Chat Completions 接口，全部返回结构化数据（Pydantic 校验）。")
bullet(doc, "成本控制：单次检测调用次数硬上限、结果缓存（相同输入复用，复跑 0 费用）、真实 token 统计。")
bullet(doc, "失败处理：超时/限流/格式错误按配置重试，仍失败则降级并在报告中声明。")

heading(doc, "（五）数据、代码与验证方法", size=12)
bullet(doc, "标注测试集 100 条，覆盖 7 类；凡涉及文献的真值均经 Crossref 实测锚定（26 个 DOI：19 个存在、7 个确认 404）。")
bullet(doc, "自动化测试 245 项，全部离线可跑、零成本；静态检查 pyflakes 0 告警。")
bullet(doc, "每条证据保留原始响应留档，可逐条回溯核对。")
bullet(doc, "代码公开于 https://github.com/depend-fy/-AI，含完整提交历史与评测存档。")

heading(doc, "三、实验结果")
body(
    doc,
    "在同一批 100 条标注用例、同一个模型的前提下，对比“直接把文本交给模型判断”（B0）"
    "与本作品完整流程（B1）：",
)
make_table(
    doc,
    ["指标", "B0 直接问模型", "B1 本项目方法"],
    [
        ["精确率", "61.4%", "97.6%"],
        ["召回率", "68.6%", "78.4%"],
        ["F1", "64.8%", "87.0%"],
        ["误报率", "44.9%", "2.0%"],
        ["六状态完全一致率", "51.0%", "79.0%"],
        ["证据可追溯率", "0.0%", "69.0%"],
        ["修正成功率 / 剩余错误率", "0.0% / 100.0%", "79.3% / 20.7%"],
        ["新增错误率", "0.0%", "0.0%"],
    ],
    widths=[136, 156, 156],
)
body(
    doc,
    "结论：B0 的错误方式是“冤枉”——把正确内容判成有问题，而不是漏检；"
    "本方法在保持低误报率的同时提升了召回率，并使多数结论可点开证据。",
)

heading(doc, "四、创新点")
bullet(doc, "把“判定需要证据”做成数据契约：违反即无法构造对象，从机制上杜绝无证据结论。")
bullet(doc, "以确定性规则表取代模型主观评分，每条结论可回溯到具体规则编号与证据。")
bullet(doc, "跨厂商独立复核修正建议，并区分“改了”与“改对了”两个问题。")
bullet(doc, "以误报率而非精确率作为核心诚实指标，并主动公开全部失败案例。")
bullet(doc, "如实区分“服务端明确不存在”与“检索失败”，后者不得用于暗示内容有问题。")

heading(doc, "五、诚实边界与局限")
bullet(doc, "只读到摘要，未通读全文；摘要层面的判断不等于全文核实。")
bullet(doc, "中文文献无法核验：Crossref / OpenAlex 未提供中文库公开接口。")
bullet(doc, "非学术事实（政策、新闻、产品参数）需付费检索服务，当前不在覆盖范围。")
bullet(doc, "作者与机构类核验未实现，目前仅实现年份与期刊比对。")
bullet(doc, "100 条测试中 21 条状态不一致（含误报、漏判与状态不精确），已全部公开。")
bullet(doc, "标注集为自建，其中 80 条为口头确认，未计算标注一致性（kappa）。")

heading(doc, "六、总结")
body(
    doc,
    "本作品把“判断一段科研文本是否可信”从模型的主观判断，转化为一套有证据、有规则、可复核的流程。"
    "其实测结果说明：直接询问模型不仅准确率有限，更关键的是结论无法追问；"
    "而引入外部证据、确定性规则与独立复核后，误报率可降至 2.0%，"
    "并使多数结论具备可追溯的证据链。"
    "后续工作包括接入开放获取全文、扩充标注集并引入独立标注者，以及扩展非学术事实的核验能力。",
)

doc.save(str(OUT / "智鉴AI-作品说明与技术方案.docx"))
print("已生成：智鉴AI-作品说明与技术方案.docx")
